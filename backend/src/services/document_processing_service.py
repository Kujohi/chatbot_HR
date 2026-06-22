from langchain_core.documents import Document
from langchain_community.document_loaders import PyMuPDFLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
import os
import logging
import tempfile
import requests
from uuid import uuid4
from dotenv import load_dotenv
from langchain_google_genai import GoogleGenerativeAIEmbeddings

from src.services.gemini_config import GEMINI_EMBEDDING_MODEL
load_dotenv()

logger = logging.getLogger(__name__)

from src.db.repositories.documents import (
    get_document_by_id,
)
from src.db.repositories.vectors import delete_vectors_by_keys, upsert_document_vectors

_embeddings = GoogleGenerativeAIEmbeddings(model=GEMINI_EMBEDDING_MODEL)


def _load_pdf(file_url: str) -> list[Document]:
    """Load a PDF from a URL using PyMuPDFLoader."""
    loader = PyMuPDFLoader(file_url)
    documents = loader.load()
    # Fix common Vietnamese PDF extraction errors (e.g. old font encodings)
    for doc in documents:
        doc.page_content = doc.page_content.replace('ƣ', 'ư').replace('Ƣ', 'Ư')
    return documents


def _load_docx(file_url: str) -> list[Document]:
    """Download a .docx from a URL and extract text using python-docx."""
    from docx import Document as DocxDocument

    response = requests.get(file_url, timeout=60)
    response.raise_for_status()

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp.write(response.content)
        tmp_path = tmp.name

    try:
        doc = DocxDocument(tmp_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        full_text = "\n".join(paragraphs)

        # Return as a single langchain Document (no page concept for docx)
        return [
            Document(
                page_content=full_text,
                metadata={"source": file_url, "page": 1},
            )
        ]
    finally:
        os.unlink(tmp_path)


def _load_document_content(file_url: str, source_type: str) -> list[Document]:
    """Dispatch to the appropriate loader based on source_type."""
    if source_type == "docx":
        return _load_docx(file_url)
    else:
        return _load_pdf(file_url)


def split_document(document_id: str):
    doc_data = get_document_by_id(document_id)

    if not doc_data:
        raise Exception(f"Document with ID {document_id} not found")

    file_url = doc_data["file_url"]
    source_type = doc_data.get("source_type", "pdf")

    documents = _load_document_content(file_url, source_type)

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1500,
        chunk_overlap=100,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks = splitter.split_documents(documents)
    
    # Inject custom metadata into each chunk
    storage_path = doc_data.get("storage_path") or ""
    for chunk in chunks:
        chunk.metadata["document_id"] = str(document_id)
        chunk.metadata["storage_path"] = storage_path
        chunk.metadata["chunk_id"] = str(uuid4())
        chunk.metadata["title"] = doc_data.get("title", "")
        chunk.metadata["owner_id"] = doc_data.get("owner_id", "")
        
    return chunks

def store_document_chunks(chunks: list[Document]):
    if not chunks:
        return

    vectors = _embeddings.embed_documents([chunk.page_content for chunk in chunks])
    insert_data = []
    for chunk, embedding in zip(chunks, vectors, strict=True):
        insert_data.append(
            {
                "document_id": chunk.metadata["document_id"],
                "chunk_id": chunk.metadata["chunk_id"],
                "vector_key": chunk.metadata["chunk_id"],
                "content": chunk.page_content,
                "metadata": chunk.metadata,
                "embedding": embedding,
                "is_document": False,
            }
        )

    upsert_document_vectors(insert_data)

def document_summary_vector_id(document_id: str) -> str:
    return f"doc-summary-{document_id}"


def index_document_summary(
    document_id: str,
    summary: str,
    storage_path: str,
    title: str = "",
    is_image_doc: bool = False,
) -> None:
    """Index a document-level summary vector for routing queries to relevant files."""
    vector_id = document_summary_vector_id(document_id)
    embedding = _embeddings.embed_query(summary)
    upsert_document_vectors(
        [
            {
                "document_id": document_id,
                "chunk_id": None,
                "vector_key": vector_id,
                "content": summary,
                "metadata": {
                    "document_id": str(document_id),
                    "isDocument": True,
                    "is_image_doc": is_image_doc,
                    "storage_path": storage_path,
                    "title": title,
                },
                "embedding": embedding,
                "is_document": True,
            }
        ]
    )
    logger.info(f"Indexed document summary for {document_id} as {vector_id}")


def index_document_vectors(chunks: list[Document]):
    store_document_chunks(chunks)

def delete_document_vectors(chunk_ids: list[str]):
    """
    Deletes document vectors from PostgreSQL.
    """
    if not chunk_ids:
        logger.info("No vector keys provided for deletion")
        return
    delete_vectors_by_keys(chunk_ids)


if __name__ == "__main__":
    print(len(split_document("17")))
